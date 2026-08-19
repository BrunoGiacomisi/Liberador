"""Genera un Excel de prueba con la misma estructura que la planilla real,
para poder probar el generador de libres sin depender de datos reales."""

from pathlib import Path

import openpyxl

RUTA_SALIDA = Path(__file__).resolve().parent / "manifiesto_prueba.xlsx"

ENCABEZADOS = [
    "POL", "POD", "BL", "T.CNTR", "TP", "QTY", "DESCRIPCION", "CHASIS",
    "CNEE", "DEPOSITO", "FLETE", "DIAS LIBRE", "FECHA LIBRE PARA RETIRO Y/O DEVOLUCION",
    "DIAS LIBRE MAFI", "FECHA LIBRE PARA EL DESCONSOLIDADO",
]

FILAS = [
    ["DHEAM", "UYMVD", "S329422175", "USED LM ROROUroro", "MH", 1, "Toyota Hilux Double CAB",
     "AH11ZG9G0914S366", "Alfred Roland Metzger", 1716, "NO", 5, "17/06/2026", "", ""],
    ["DHEAM", "UYMVD", "S329169920", "USED SMALL VAN(S)/TUR", "MH", 1, "Westfalia-California",
     "WV2ZZZ70ZWH095793", "Arnold Andreas Dr. Mager", 1716, "NO", 5, "17/06/2026", "", ""],
    ["ESIBO", "BLANR", "S329400037", "NEW SMALL VAN(S)/VEH", "VEH", 1, "MERCEDES BENZ VITOMERCEDES",
     "W1WV1G9F0794026503", "Autoblicor Uruguay S.A.", "BOMPORT", "NO", "", "", "", ""],
    ["BLANR", "BLANR", "S329612280", "NEW CAR(S)/VNRoro", "VEH", 1, "BMW X8BMW",
     "WBA11G9F0794026503", "Automotres Motor Haus S.A.", "BOMPORT", "NO", "", "", "", ""],
    ["DHEAM", "UYMVD", "S329568298", "USED LM ROROUroro", "VEH", 1, "PORSCHE MACAN GTSPORSCHE",
     "WP1ZZZXAZTL251414", "Nordenwagen Uruguay Sa", 1716, "NO", 5, "17/06/2026", "", ""],
    ["BRPNS", "BRPNS", "S329509809", "VEH", "VEH", 503, "RENAULT",
     "PENDING", "SANTA ROSA AUTOMOTORS", "DD1651", "SI", "", "", "", ""],
    ["DHEAM", "UYMVD", "S329701796", "40HC", "40HC", 1, "Consolidated Shipments stc 55 pkgs spare parts",
     "GCNU4865094", "Jauser Soluciones Logisticas", 1714, "NO", 5, "28/08/2026", "", ""],
    ["DHEAM", "UYMVD", "S329701796", "40HC", "40HC", 1, "Consolidated Shipments stc 44 pkgs spare parts",
     "ACLU9803492", "Jauser Soluciones Logisticas", 1714, "NO", 5, "28/08/2026", "", ""],
]

TITULO_FILA1 = "PGD GRANDE FRANCIA 0426 ETA 4/7/25 ESCALA:  MANIFIESTO: 198401"


def main() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hoja1"

    ws.cell(row=1, column=1, value=TITULO_FILA1)
    for col_idx, encabezado in enumerate(ENCABEZADOS, start=1):
        ws.cell(row=2, column=col_idx, value=encabezado)

    for fila_idx, fila in enumerate(FILAS, start=3):
        for col_idx, valor in enumerate(fila, start=1):
            ws.cell(row=fila_idx, column=col_idx, value=valor)

    wb.save(RUTA_SALIDA)
    print(f"Excel de prueba creado en: {RUTA_SALIDA}")


if __name__ == "__main__":
    main()
